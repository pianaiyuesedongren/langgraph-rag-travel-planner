from __future__ import annotations

from pathlib import Path

from langchain_core.documents import Document


class DocumentLoader:
    """文档加载器 - 支持PDF/Word/Markdown/TXT"""

    SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".doc", ".md", ".txt", ".json"}

    def __init__(self, directory: str | Path):
        self.directory = Path(directory)

    def load(self) -> list[Document]:
        """加载目录下所有支持的文档"""
        documents = []
        if not self.directory.exists():
            return documents

        for file_path in self.directory.rglob("*"):
            if file_path.suffix.lower() in self.SUPPORTED_EXTENSIONS:
                docs = self._load_file(file_path)
                documents.extend(docs)

        return documents

    def _load_file(self, file_path: Path) -> list[Document]:
        """加载单个文件"""
        suffix = file_path.suffix.lower()

        try:
            if suffix == ".pdf":
                return self._load_pdf(file_path)
            elif suffix in {".docx", ".doc"}:
                return self._load_docx(file_path)
            elif suffix == ".md":
                return self._load_markdown(file_path)
            elif suffix == ".txt":
                return self._load_text(file_path)
            elif suffix == ".json":
                return self._load_json(file_path)
        except Exception as e:
            print(f"Failed to load {file_path}: {e}")
            return []

        return []

    def _load_pdf(self, file_path: Path) -> list[Document]:
        """加载PDF文件"""
        from pypdf import PdfReader

        reader = PdfReader(str(file_path))
        documents = []

        for i, page in enumerate(reader.pages):
            text = page.extract_text()
            if text and text.strip():
                documents.append(
                    Document(
                        page_content=text.strip(),
                        metadata={
                            "source": str(file_path),
                            "page": i + 1,
                            "type": "pdf",
                        },
                    )
                )

        return documents

    def _load_docx(self, file_path: Path) -> list[Document]:
        """加载Word文档"""
        from docx import Document as DocxDocument

        doc = DocxDocument(str(file_path))
        documents = []

        for i, para in enumerate(doc.paragraphs):
            if para.text and para.text.strip():
                documents.append(
                    Document(
                        page_content=para.text.strip(),
                        metadata={
                            "source": str(file_path),
                            "paragraph": i,
                            "type": "docx",
                        },
                    )
                )

        return documents

    def _load_markdown(self, file_path: Path) -> list[Document]:
        """加载Markdown文件"""
        content = file_path.read_text(encoding="utf-8")
        if not content.strip():
            return []

        return [
            Document(
                page_content=content.strip(),
                metadata={
                    "source": str(file_path),
                    "type": "markdown",
                },
            )
        ]

    def _load_text(self, file_path: Path) -> list[Document]:
        """加载纯文本文件"""
        content = file_path.read_text(encoding="utf-8")
        if not content.strip():
            return []

        return [
            Document(
                page_content=content.strip(),
                metadata={
                    "source": str(file_path),
                    "type": "text",
                },
            )
        ]

    def _load_json(self, file_path: Path) -> list[Document]:
        """加载JSON文件"""
        import json

        content = file_path.read_text(encoding="utf-8")
        data = json.loads(content)
        documents = []

        if isinstance(data, list):
            for i, item in enumerate(data):
                text = json.dumps(item, ensure_ascii=False, indent=2)
                documents.append(
                    Document(
                        page_content=text,
                        metadata={
                            "source": str(file_path),
                            "index": i,
                            "type": "json",
                        },
                    )
                )
        elif isinstance(data, dict):
            text = json.dumps(data, ensure_ascii=False, indent=2)
            documents.append(
                Document(
                    page_content=text,
                    metadata={
                        "source": str(file_path),
                        "type": "json",
                    },
                )
            )

        return documents
